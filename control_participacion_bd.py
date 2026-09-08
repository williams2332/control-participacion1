"""Control de Participación: Tkinter + SQLite + Excel.

Características:
- Usa SQLite como base de datos principal.
- Sincroniza una copia legible en participacion.xlsx.
- Muestra el promedio general de participación.
- Colorea el promedio y las filas: rojo, amarillo o verde.
- Reproduce un sonido después de guardar correctamente.
- Carga el logo desde la ruta indicada por el usuario.

Dependencias:
    pip install openpyxl pillow
"""

from __future__ import annotations

import os
import sqlite3
import subprocess
import sys
from datetime import date, datetime
from pathlib import Path
import tkinter as tk
from tkinter import messagebox, ttk

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

try:
    from PIL import Image, ImageTk
except ImportError:
    Image = None
    ImageTk = None


NOMBRE_APP = "Control de Participación"
NOMBRE_HOJA = "Participación"
NOMBRE_ARCHIVO_EXCEL = "participacion.xlsx"
NOMBRE_ARCHIVO_BD = "participacion.db"

# Esta es la ruta solicitada. También se prueban rutas alternativas para
# que el programa no falle si se copia a otra computadora.
RUTA_LOGO_USUARIO = Path(r"C:\Users\willi\Desktop\Paticipacion\Logo.jpg")

ENCABEZADOS_EXCEL = (
    "Fecha",
    "Curso",
    "Número de lista",
    "Nombre",
    "Participación",
    "Observación",
)

TIPOS_PARTICIPACION = (
    "Excelente",
    "Buena",
    "Regular",
    "Pregunta",
    "Respuesta",
    "Exposición",
    "Debate",
    "Trabajo en grupo",
    "Otra",
)

# Escala utilizada para calcular el promedio general.
PUNTUACIONES = {
    "Excelente": 4.0,
    "Buena": 3.0,
    "Regular": 2.0,
    "Pregunta": 3.0,
    "Respuesta": 3.0,
    "Exposición": 4.0,
    "Debate": 3.0,
    "Trabajo en grupo": 3.0,
    "Otra": 2.0,
}

COLOR_FONDO = "#F4F7FB"
COLOR_PRIMARIO = "#1F5A94"
COLOR_PRIMARIO_OSCURO = "#16436E"
COLOR_TEXTO = "#1F2937"
COLOR_BORDE = "#D6E0EC"
COLOR_ROJO = "#FDE2E2"
COLOR_ROJO_TEXTO = "#A61B1B"
COLOR_AMARILLO = "#FFF4CC"
COLOR_AMARILLO_TEXTO = "#8A5A00"
COLOR_VERDE = "#DDF5E4"
COLOR_VERDE_TEXTO = "#146B32"


def ruta_recurso(nombre: str) -> Path:
    """
    Devuelve la ruta de un recurso tanto en modo Python como dentro
    de un ejecutable creado con PyInstaller.
    """
    carpeta_empaquetada = getattr(sys, "_MEIPASS", None)

    if carpeta_empaquetada:
        return Path(carpeta_empaquetada) / nombre

    return Path(__file__).resolve().parent / nombre


class ControlParticipacionApp:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title(NOMBRE_APP)
        self.root.geometry("1280x820")
        self.root.minsize(1000, 650)
        self.root.configure(bg=COLOR_FONDO)

        carpeta_documentos = self.obtener_carpeta_documentos()
        self.carpeta_datos = carpeta_documentos / "Control de Participacion"
        self.carpeta_datos.mkdir(parents=True, exist_ok=True)
        self.archivo_bd = self.carpeta_datos / NOMBRE_ARCHIVO_BD
        self.archivo_excel = self.carpeta_datos / NOMBRE_ARCHIVO_EXCEL
        self.ruta_logo = self.buscar_logo()

        self.conexion: sqlite3.Connection | None = None
        self.logo_imagen = None
        self.registros: list[dict] = []

        self.var_busqueda = tk.StringVar()
        self.var_estado = tk.StringVar(value="Listo para registrar una participación")
        self.var_promedio = tk.StringVar(value="Promedio: —")
        self.var_nombre = tk.StringVar()
        self.var_numero = tk.StringVar()
        self.var_curso = tk.StringVar()
        self.var_fecha = tk.StringVar(value=self.fecha_actual())
        self.var_participacion = tk.StringVar()
        self.var_observacion = tk.StringVar()

        self.configurar_estilos()
        self.inicializar_base_datos()
        self.migrar_excel_antiguo_si_corresponde()
        self.sincronizar_excel()
        self.crear_interfaz()
        self.cargar_registros()

        self.entrada_nombre.focus_set()
        self.root.bind("<Control-s>", self.registrar_desde_teclado)
        self.root.bind("<Control-f>", self.enfocar_busqueda)
        self.root.bind("<Delete>", self.eliminar_seleccionado_desde_teclado)
        self.root.protocol("WM_DELETE_WINDOW", self.cerrar)

    @staticmethod
    def obtener_carpeta_documentos() -> Path:
        candidatos = (Path.home() / "Documents", Path.home() / "Documentos")
        for carpeta in candidatos:
            if carpeta.exists() or carpeta.parent.exists():
                carpeta.mkdir(parents=True, exist_ok=True)
                return carpeta
        return Path.home()

    def buscar_logo(self) -> Path:
        candidatos = (
            ruta_recurso("Logo.jpg"),
            ruta_recurso("logo.jpg"),
            RUTA_LOGO_USUARIO,
            Path.home() / "Desktop" / "Paticipacion" / "Logo.jpg",
            Path.home() / "OneDrive" / "Desktop" / "Paticipacion" / "Logo.jpg",
        )
        for candidato in candidatos:
            if candidato.exists():
                return candidato
        # Se conserva la ruta solicitada para que el mensaje de error sea claro.
        return RUTA_LOGO_USUARIO

    @staticmethod
    def fecha_actual() -> str:
        return datetime.now().strftime("%d/%m/%Y")

    def configurar_estilos(self) -> None:
        estilo = ttk.Style(self.root)
        try:
            estilo.theme_use("clam")
        except tk.TclError:
            pass

        estilo.configure("TFrame", background=COLOR_FONDO)
        estilo.configure("TLabel", background=COLOR_FONDO, foreground=COLOR_TEXTO,
                         font=("Segoe UI", 10))
        estilo.configure("TButton", font=("Segoe UI", 10), padding=(12, 7))
        estilo.configure("Primary.TButton", background=COLOR_PRIMARIO,
                         foreground="white", font=("Segoe UI", 10, "bold"),
                         padding=(14, 8))
        estilo.map("Primary.TButton", background=[("active", COLOR_PRIMARIO_OSCURO)])
        estilo.configure("Title.TLabel", background=COLOR_FONDO,
                         foreground=COLOR_PRIMARIO_OSCURO,
                         font=("Segoe UI", 24, "bold"))
        estilo.configure("Subtitle.TLabel", background=COLOR_FONDO,
                         foreground="#526579", font=("Segoe UI", 11))
        estilo.configure("TLabelframe", background=COLOR_FONDO,
                         bordercolor=COLOR_BORDE)
        estilo.configure("TLabelframe.Label", background=COLOR_FONDO,
                         foreground=COLOR_PRIMARIO_OSCURO,
                         font=("Segoe UI", 10, "bold"))
        estilo.configure("Treeview", rowheight=29, font=("Segoe UI", 10),
                         background="white", fieldbackground="white",
                         foreground=COLOR_TEXTO)
        estilo.configure("Treeview.Heading", font=("Segoe UI", 10, "bold"),
                         background=COLOR_PRIMARIO, foreground="white", padding=7)
        estilo.map("Treeview", background=[("selected", "#CFE3F7")],
                   foreground=[("selected", COLOR_TEXTO)])
        estilo.configure("Status.TLabel", background="#E7EEF7",
                         foreground="#425466", padding=(10, 7))

    def inicializar_base_datos(self) -> None:
        self.conexion = sqlite3.connect(self.archivo_bd)
        self.conexion.row_factory = sqlite3.Row
        self.conexion.execute("PRAGMA foreign_keys = ON")
        self.conexion.execute(
            """
            CREATE TABLE IF NOT EXISTS participaciones (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                fecha TEXT NOT NULL,
                curso TEXT NOT NULL,
                numero_lista INTEGER NOT NULL,
                nombre TEXT NOT NULL,
                participacion TEXT NOT NULL,
                observacion TEXT DEFAULT '',
                puntuacion REAL NOT NULL,
                creado_en TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
            """
        )
        self.conexion.execute(
            """
            CREATE TABLE IF NOT EXISTS configuracion (
                clave TEXT PRIMARY KEY,
                valor TEXT NOT NULL
            )
            """
        )
        self.conexion.commit()

    def migrar_excel_antiguo_si_corresponde(self) -> None:
        """Importa una sola vez los registros del Excel creado por la versión anterior."""
        if self.conexion is None:
            return

        migrado = self.conexion.execute(
            "SELECT valor FROM configuracion WHERE clave = 'excel_migrado'"
        ).fetchone()
        if migrado:
            return

        if self.archivo_excel.exists():
            try:
                from openpyxl import load_workbook

                libro = load_workbook(self.archivo_excel, data_only=True, read_only=True)
                if NOMBRE_HOJA in libro.sheetnames:
                    hoja = libro[NOMBRE_HOJA]
                    for fila in hoja.iter_rows(min_row=2, max_col=6, values_only=True):
                        if not any(valor not in (None, "") for valor in fila):
                            continue
                        fecha = self.normalizar_fecha(fila[0])
                        curso = str(fila[1] or "Sin curso")
                        numero = self.convertir_numero(fila[2])
                        nombre = str(fila[3] or "Sin nombre")
                        participacion = str(fila[4] or "Otra")
                        if participacion not in TIPOS_PARTICIPACION:
                            participacion = "Otra"
                        observacion = str(fila[5] or "")
                        self.insertar_registro(
                            fecha, curso, numero, nombre, participacion, observacion,
                            confirmar=False,
                        )
                libro.close()
            except Exception as error:
                # La base se crea aun si el Excel antiguo no puede leerse.
                self.var_estado.set(f"No se pudo migrar el Excel antiguo: {error}")

        self.conexion.execute(
            "INSERT OR REPLACE INTO configuracion (clave, valor) VALUES ('excel_migrado', '1')"
        )
        self.conexion.commit()

    @staticmethod
    def normalizar_fecha(valor) -> str:
        if isinstance(valor, (datetime, date)):
            return valor.strftime("%Y-%m-%d")
        texto = str(valor or "").strip()
        for formato in ("%d/%m/%Y", "%Y-%m-%d", "%m/%d/%Y"):
            try:
                return datetime.strptime(texto, formato).strftime("%Y-%m-%d")
            except ValueError:
                continue
        return datetime.now().strftime("%Y-%m-%d")

    @staticmethod
    def convertir_numero(valor) -> int:
        try:
            numero = int(float(valor))
            return max(1, numero)
        except (TypeError, ValueError):
            return 1

    def insertar_registro(
        self, fecha: str, curso: str, numero: int, nombre: str,
        participacion: str, observacion: str, confirmar: bool = True,
    ) -> None:
        if self.conexion is None:
            raise RuntimeError("La base de datos no está disponible.")
        puntuacion = PUNTUACIONES.get(participacion, 2.0)
        self.conexion.execute(
            """
            INSERT INTO participaciones
            (fecha, curso, numero_lista, nombre, participacion, observacion, puntuacion)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (fecha, curso, numero, nombre, participacion, observacion, puntuacion),
        )
        if confirmar:
            self.conexion.commit()

    def sincronizar_excel(self) -> None:
        """Genera el Excel desde SQLite para que ambos formatos estén alineados."""
        if self.conexion is None:
            return

        libro = Workbook()
        hoja = libro.active
        hoja.title = NOMBRE_HOJA
        hoja.append(ENCABEZADOS_EXCEL)

        filas = self.conexion.execute(
            """
            SELECT fecha, curso, numero_lista, nombre, participacion, observacion
            FROM participaciones
            ORDER BY id
            """
        ).fetchall()
        for fila in filas:
            fecha = datetime.strptime(fila["fecha"], "%Y-%m-%d")
            hoja.append([
                fecha, fila["curso"], fila["numero_lista"], fila["nombre"],
                fila["participacion"], fila["observacion"],
            ])
            hoja.cell(row=hoja.max_row, column=1).number_format = "DD/MM/YYYY"

        for celda in hoja[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="1F5A94")
            celda.alignment = Alignment(horizontal="center", vertical="center")

        for fila in hoja.iter_rows(min_row=2):
            for celda in fila:
                celda.alignment = Alignment(vertical="center", wrap_text=True)

        anchos = (15, 20, 17, 32, 22, 45)
        for indice, ancho in enumerate(anchos, start=1):
            hoja.column_dimensions[hoja.cell(row=1, column=indice).column_letter].width = ancho
        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = f"A1:F{max(1, hoja.max_row)}"
        hoja.row_dimensions[1].height = 24

        if filas:
            tabla = Table(displayName="TablaParticipacion", ref=f"A1:F{hoja.max_row}")
            tabla.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2", showFirstColumn=False,
                showLastColumn=False, showRowStripes=True, showColumnStripes=False,
            )
            hoja.add_table(tabla)

        libro.save(self.archivo_excel)

    def crear_interfaz(self) -> None:
        encabezado = ttk.Frame(self.root)
        encabezado.pack(fill="x", padx=28, pady=(20, 8))

        bloque_titulo = ttk.Frame(encabezado)
        bloque_titulo.pack(side="left", fill="x", expand=True)
        ttk.Label(bloque_titulo, text="AULA · CONTROL DE PARTICIPACIÓN",
                  style="Title.TLabel").pack(anchor="w")
        ttk.Label(
            bloque_titulo,
            text="Registra, consulta y mide la participación de tus estudiantes",
            style="Subtitle.TLabel",
        ).pack(anchor="w", pady=(3, 0))

        resumen = tk.Frame(encabezado, bg="white", highlightbackground=COLOR_BORDE,
                           highlightthickness=1, padx=14, pady=8)
        resumen.pack(side="right", padx=(10, 16))
        tk.Label(resumen, textvariable=self.var_promedio, bg="white",
                 fg=COLOR_PRIMARIO_OSCURO, font=("Segoe UI", 15, "bold")).pack()
        self.etiqueta_estado_promedio = tk.Label(
            resumen, text="Sin registros", bg="white", fg="#667085",
            font=("Segoe UI", 10, "bold"),
        )
        self.etiqueta_estado_promedio.pack(pady=(2, 0))

        self.cargar_logo(encabezado)
        self.crear_formulario()
        self.crear_barra_acciones()
        self.crear_tabla()
        ttk.Label(self.root, textvariable=self.var_estado,
                  style="Status.TLabel", anchor="w").pack(fill="x", side="bottom")

    def cargar_logo(self, contenedor) -> None:
        if Image is None or ImageTk is None:
            tk.Label(
                contenedor,
                text="Falta Pillow\npip install pillow",
                bg=COLOR_FONDO,
                fg=COLOR_ROJO_TEXTO,
                font=("Segoe UI", 9, "bold"),
                justify="center",
            ).pack(side="right", padx=(10, 0))
            self.var_estado.set(
                "No se cargó el logo: instala Pillow con "
                "'pip install pillow'."
            )
            return

        if not self.ruta_logo.exists():
            tk.Label(
                contenedor,
                text="Logo no encontrado",
                bg=COLOR_FONDO,
                fg=COLOR_ROJO_TEXTO,
                font=("Segoe UI", 9, "bold"),
            ).pack(side="right", padx=(10, 0))
            self.var_estado.set(
                "No se encontró Logo.jpg. Colócalo junto al programa "
                "o en C:\\Users\\willi\\Desktop\\Paticipacion\\."
            )
            return

        try:
            imagen = Image.open(self.ruta_logo)
            imagen.thumbnail((100, 100))
            self.logo_imagen = ImageTk.PhotoImage(imagen)
            tk.Label(contenedor, image=self.logo_imagen, bg=COLOR_FONDO).pack(side="right")
        except Exception as error:
            self.var_estado.set(f"No se pudo cargar el logo: {error}")

    def crear_formulario(self) -> None:
        datos = ttk.LabelFrame(self.root, text=" Datos del estudiante ", padding=14)
        datos.pack(fill="x", padx=28, pady=(6, 8))
        datos.columnconfigure(1, weight=1)
        datos.columnconfigure(3, weight=1)

        self.agregar_campo(datos, "Nombre:", self.var_nombre, 0, 0, 35)
        self.agregar_campo(datos, "Número de lista:", self.var_numero, 0, 2, 15)
        self.agregar_campo(datos, "Curso:", self.var_curso, 1, 0, 35)
        self.agregar_campo(datos, "Fecha (dd/mm/aaaa):", self.var_fecha, 1, 2, 15)

        participacion = ttk.LabelFrame(self.root, text=" Participación ", padding=14)
        participacion.pack(fill="x", padx=28, pady=(0, 8))
        participacion.columnconfigure(1, weight=1)
        participacion.columnconfigure(3, weight=1)

        ttk.Label(participacion, text="Tipo de participación:").grid(
            row=0, column=0, sticky="w", padx=(2, 8), pady=6)
        self.combo_participacion = ttk.Combobox(
            participacion, textvariable=self.var_participacion,
            values=TIPOS_PARTICIPACION, state="readonly", width=28,
        )
        self.combo_participacion.grid(row=0, column=1, sticky="ew", padx=8, pady=6)

        ttk.Label(participacion, text="Observación:").grid(
            row=0, column=2, sticky="w", padx=(16, 8), pady=6)
        self.entrada_observacion = ttk.Entry(
            participacion, textvariable=self.var_observacion, width=45,
        )
        self.entrada_observacion.grid(row=0, column=3, sticky="ew", padx=8, pady=6)

    def agregar_campo(self, contenedor, etiqueta, variable, fila, columna, ancho) -> None:
        ttk.Label(contenedor, text=etiqueta).grid(
            row=fila, column=columna, sticky="w", padx=(2, 8), pady=6)
        entrada = ttk.Entry(contenedor, textvariable=variable, width=ancho)
        entrada.grid(row=fila, column=columna + 1, sticky="ew", padx=8, pady=6)

        if variable is self.var_nombre:
            self.entrada_nombre = entrada
        elif variable is self.var_numero:
            self.entrada_numero = entrada
        elif variable is self.var_curso:
            self.entrada_curso = entrada
        elif variable is self.var_fecha:
            self.entrada_fecha = entrada

    def crear_barra_acciones(self) -> None:
        acciones = ttk.Frame(self.root)
        acciones.pack(fill="x", padx=28, pady=(0, 10))
        ttk.Button(acciones, text="Registrar participación",
                   style="Primary.TButton", command=self.registrar).pack(side="left", padx=(0, 7))
        ttk.Button(acciones, text="Limpiar formulario", command=self.limpiar).pack(side="left", padx=4)
        ttk.Button(acciones, text="Eliminar seleccionado",
                   command=self.eliminar_seleccionado).pack(side="left", padx=4)
        ttk.Button(acciones, text="Abrir Excel", command=self.abrir_excel).pack(side="right", padx=(7, 0))
        ttk.Button(acciones, text="Abrir carpeta", command=self.abrir_carpeta).pack(side="right", padx=4)

        busqueda = ttk.Frame(self.root)
        busqueda.pack(fill="x", padx=28, pady=(0, 8))
        ttk.Label(busqueda, text="Buscar:").pack(side="left", padx=(0, 8))
        self.entrada_busqueda = ttk.Entry(busqueda, textvariable=self.var_busqueda, width=44)
        self.entrada_busqueda.pack(side="left")
        ttk.Label(busqueda, text="  Nombre, curso, tipo u observación",
                  foreground="#68778A").pack(side="left")
        self.var_busqueda.trace_add("write", lambda *_: self.mostrar_registros())
        self.crear_leyenda_colores()

    def crear_leyenda_colores(self) -> None:
        leyenda = ttk.Frame(self.root)
        leyenda.pack(fill="x", padx=28, pady=(0, 8))
        ttk.Label(leyenda, text="Indicador:").pack(side="left", padx=(0, 8))
        self.crear_indicador(leyenda, "Baja (< 2.5)", COLOR_ROJO, COLOR_ROJO_TEXTO)
        self.crear_indicador(leyenda, "Regular (2.5–3.49)", COLOR_AMARILLO, COLOR_AMARILLO_TEXTO)
        self.crear_indicador(leyenda, "Buena (≥ 3.5)", COLOR_VERDE, COLOR_VERDE_TEXTO)

    @staticmethod
    def crear_indicador(contenedor, texto: str, fondo: str, color: str) -> None:
        tk.Label(
            contenedor,
            text=f"  {texto}  ",
            bg=fondo,
            fg=color,
            font=("Segoe UI", 9, "bold"),
            padx=4,
            pady=2,
        ).pack(side="left", padx=4)

    def crear_tabla(self) -> None:
        contenedor = ttk.LabelFrame(self.root, text=" Registros de participación ", padding=10)
        contenedor.pack(fill="both", expand=True, padx=28, pady=(0, 18))
        contenedor.columnconfigure(0, weight=1)
        contenedor.rowconfigure(0, weight=1)

        columnas = ("fecha", "curso", "numero", "nombre", "participacion", "observacion")
        self.tabla = ttk.Treeview(contenedor, columns=columnas, show="headings", selectmode="extended")
        encabezados = {
            "fecha": "Fecha", "curso": "Curso", "numero": "No. lista",
            "nombre": "Nombre", "participacion": "Participación", "observacion": "Observación",
        }
        anchos = {
            "fecha": (105, "center"), "curso": (130, "center"), "numero": (95, "center"),
            "nombre": (230, "w"), "participacion": (165, "center"), "observacion": (350, "w"),
        }
        for columna in columnas:
            self.tabla.heading(columna, text=encabezados[columna])
            ancho, alineacion = anchos[columna]
            self.tabla.column(columna, width=ancho, anchor=alineacion, minwidth=70)

        self.tabla.tag_configure("rojo", background=COLOR_ROJO, foreground=COLOR_ROJO_TEXTO)
        self.tabla.tag_configure("amarillo", background=COLOR_AMARILLO, foreground=COLOR_AMARILLO_TEXTO)
        self.tabla.tag_configure("verde", background=COLOR_VERDE, foreground=COLOR_VERDE_TEXTO)

        scroll_vertical = ttk.Scrollbar(contenedor, orient="vertical", command=self.tabla.yview)
        scroll_horizontal = ttk.Scrollbar(contenedor, orient="horizontal", command=self.tabla.xview)
        self.tabla.configure(yscrollcommand=scroll_vertical.set, xscrollcommand=scroll_horizontal.set)
        self.tabla.grid(row=0, column=0, sticky="nsew")
        scroll_vertical.grid(row=0, column=1, sticky="ns")
        scroll_horizontal.grid(row=1, column=0, sticky="ew")
        self.tabla.bind("<Double-1>", self.cargar_seleccion_en_formulario)

    def datos_del_formulario(self) -> tuple[str, int, str, str, str, str] | None:
        nombre = self.var_nombre.get().strip()
        numero_texto = self.var_numero.get().strip()
        curso = self.var_curso.get().strip()
        fecha_texto = self.var_fecha.get().strip()
        participacion = self.var_participacion.get().strip()
        observacion = self.var_observacion.get().strip()

        if not nombre:
            return self.advertir("Escribe el nombre del estudiante.", self.entrada_nombre)
        if not curso:
            return self.advertir("Escribe el curso.", self.entrada_curso)
        if not numero_texto.isdigit() or int(numero_texto) <= 0:
            return self.advertir("El número de lista debe ser un entero positivo.", self.entrada_numero)
        try:
            fecha = datetime.strptime(fecha_texto, "%d/%m/%Y").strftime("%Y-%m-%d")
        except ValueError:
            return self.advertir(
                "La fecha debe tener el formato dd/mm/aaaa, por ejemplo 08/09/2026.",
                self.entrada_fecha,
            )
        if not participacion:
            return self.advertir("Selecciona el tipo de participación.", self.combo_participacion)

        return nombre, int(numero_texto), curso, fecha, participacion, observacion

    @staticmethod
    def advertir(mensaje: str, widget):
        messagebox.showwarning("Datos incompletos", mensaje)
        widget.focus_set()
        return None

    def registrar(self) -> None:
        datos = self.datos_del_formulario()
        if datos is None or self.conexion is None:
            return

        nombre, numero, curso, fecha, participacion, observacion = datos
        try:
            self.insertar_registro(fecha, curso, numero, nombre, participacion, observacion)
            self.conexion.commit()
            # El sonido solo ocurre después de guardar correctamente en SQLite.
            self.reproducir_sonido()
            try:
                self.sincronizar_excel()
            except PermissionError:
                messagebox.showwarning(
                    "Registro guardado",
                    "La participación se guardó en la base de datos, pero Excel está abierto. "
                    "Ciérralo y pulsa Abrir Excel para actualizarlo.",
                )
            self.cargar_registros()
        except Exception as error:
            self.conexion.rollback()
            messagebox.showerror("Error al guardar", f"No se pudo guardar el registro:\n\n{error}")
            return

        self.var_nombre.set("")
        self.var_numero.set("")
        self.var_observacion.set("")
        self.var_participacion.set("")
        self.entrada_nombre.focus_set()
        self.var_estado.set(f"Participación guardada. Total: {len(self.registros)}")

    def reproducir_sonido(self) -> None:
        try:
            if sys.platform == "win32":
                import winsound
                winsound.MessageBeep(winsound.MB_OK)
            else:
                self.root.bell()
        except Exception:
            # El guardado no debe fallar si el equipo no tiene sonido disponible.
            pass

    def cargar_registros(self) -> None:
        if self.conexion is None:
            return
        filas = self.conexion.execute(
            """
            SELECT id, fecha, curso, numero_lista, nombre, participacion, observacion, puntuacion
            FROM participaciones ORDER BY id DESC
            """
        ).fetchall()
        self.registros = [dict(fila) for fila in filas]
        self.mostrar_registros()
        self.actualizar_promedio()

    @staticmethod
    def categoria_color(puntuacion: float) -> str:
        if puntuacion < 2.5:
            return "rojo"
        if puntuacion < 3.5:
            return "amarillo"
        return "verde"

    def mostrar_registros(self) -> None:
        for elemento in self.tabla.get_children():
            self.tabla.delete(elemento)

        filtro = self.var_busqueda.get().strip().casefold()
        visibles = 0
        for registro in self.registros:
            valores = (
                datetime.strptime(registro["fecha"], "%Y-%m-%d").strftime("%d/%m/%Y"),
                registro["curso"], registro["numero_lista"], registro["nombre"],
                registro["participacion"], registro["observacion"],
            )
            texto_busqueda = " ".join(str(valor) for valor in valores).casefold()
            if filtro and filtro not in texto_busqueda:
                continue
            tag = self.categoria_color(float(registro["puntuacion"]))
            self.tabla.insert("", "end", iid=str(registro["id"]), values=valores, tags=(tag,))
            visibles += 1

        if filtro:
            self.var_estado.set(f"Mostrando {visibles} de {len(self.registros)} registro(s)")
        else:
            self.var_estado.set(f"{len(self.registros)} registro(s) cargado(s)")

    def actualizar_promedio(self) -> None:
        if self.conexion is None:
            return
        resultado = self.conexion.execute(
            "SELECT COUNT(*) AS total, AVG(puntuacion) AS promedio FROM participaciones"
        ).fetchone()
        total = int(resultado["total"])
        promedio = resultado["promedio"]
        if total == 0 or promedio is None:
            self.var_promedio.set("Promedio: —")
            self.etiqueta_estado_promedio.configure(text="Sin registros", fg="#667085")
            return

        promedio = float(promedio)
        categoria = self.categoria_color(promedio)
        textos = {"rojo": "Participación baja", "amarillo": "Participación regular",
                  "verde": "Buena participación"}
        colores = {"rojo": COLOR_ROJO_TEXTO, "amarillo": COLOR_AMARILLO_TEXTO,
                   "verde": COLOR_VERDE_TEXTO}
        self.var_promedio.set(f"Promedio: {promedio:.2f} / 4.00")
        self.etiqueta_estado_promedio.configure(text=textos[categoria], fg=colores[categoria])

    def limpiar(self) -> None:
        self.var_nombre.set("")
        self.var_numero.set("")
        self.var_curso.set("")
        self.var_fecha.set(self.fecha_actual())
        self.var_participacion.set("")
        self.var_observacion.set("")
        self.var_busqueda.set("")
        self.entrada_nombre.focus_set()
        self.var_estado.set("Formulario limpio")

    def eliminar_seleccionado(self) -> None:
        if self.conexion is None:
            return
        seleccion = self.tabla.selection()
        if not seleccion:
            messagebox.showinfo("Selecciona un registro", "Selecciona al menos un registro para eliminar.")
            return

        cantidad = len(seleccion)
        if not messagebox.askyesno(
            "Confirmar eliminación",
            f"¿Eliminar {cantidad} registro(s)?\n\nEsta acción no se puede deshacer.",
            icon="warning",
        ):
            return

        try:
            ids = [int(item) for item in seleccion]
            marcadores = ",".join("?" for _ in ids)
            self.conexion.execute(f"DELETE FROM participaciones WHERE id IN ({marcadores})", ids)
            self.conexion.commit()
            try:
                self.sincronizar_excel()
            except PermissionError:
                messagebox.showwarning(
                    "Eliminación guardada",
                    "Se eliminó de la base de datos, pero Excel está abierto y no pudo actualizarse.",
                )
            self.cargar_registros()
            self.var_estado.set(f"{cantidad} registro(s) eliminado(s)")
        except Exception as error:
            self.conexion.rollback()
            messagebox.showerror("Error al eliminar", f"No se pudieron eliminar los registros:\n\n{error}")

    def cargar_seleccion_en_formulario(self, _evento=None) -> None:
        seleccion = self.tabla.selection()
        if not seleccion:
            return
        registro = next((r for r in self.registros if str(r["id"]) == seleccion[0]), None)
        if not registro:
            return
        self.var_fecha.set(datetime.strptime(registro["fecha"], "%Y-%m-%d").strftime("%d/%m/%Y"))
        self.var_curso.set(registro["curso"])
        self.var_numero.set(str(registro["numero_lista"]))
        self.var_nombre.set(registro["nombre"])
        self.var_participacion.set(registro["participacion"])
        self.var_observacion.set(registro["observacion"])
        self.var_estado.set("Registro cargado en el formulario")

    def abrir_en_sistema(self, ruta: Path) -> None:
        ruta = Path(ruta).resolve()
        if sys.platform == "win32":
            os.startfile(str(ruta))
        elif sys.platform == "darwin":
            subprocess.Popen(["open", str(ruta)])
        else:
            subprocess.Popen(["xdg-open", str(ruta)])

    def abrir_carpeta(self) -> None:
        try:
            self.abrir_en_sistema(self.carpeta_datos)
        except Exception:
            messagebox.showinfo("Ubicación de datos", str(self.carpeta_datos))

    def abrir_excel(self) -> None:
        try:
            self.sincronizar_excel()
            self.abrir_en_sistema(self.archivo_excel)
        except Exception:
            messagebox.showinfo("Ubicación del archivo", str(self.archivo_excel))

    def registrar_desde_teclado(self, _evento=None):
        self.registrar()
        return "break"

    def eliminar_seleccionado_desde_teclado(self, _evento=None):
        if self.tabla.selection():
            self.eliminar_seleccionado()
            return "break"
        return None

    def enfocar_busqueda(self, _evento=None):
        self.entrada_busqueda.focus_set()
        self.entrada_busqueda.select_range(0, tk.END)
        return "break"

    def cerrar(self) -> None:
        if self.conexion is not None:
            self.conexion.close()
        self.root.destroy()


def main() -> None:
    try:
        root = tk.Tk()
        ControlParticipacionApp(root)
        root.mainloop()
    except Exception as error:
        try:
            messagebox.showerror(NOMBRE_APP, str(error))
        except tk.TclError:
            print(f"Error: {error}")


if __name__ == "__main__":
    main()
